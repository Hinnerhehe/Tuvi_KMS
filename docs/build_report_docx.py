# -*- coding: utf-8 -*-
"""Build docs/REPORT_KMS_TU_VI.docx from the Markdown report.

Produces a formal Vietnamese academic-report layout:
  - cover page (bìa)
  - abstract (Tóm tắt)
  - table of contents (MỤC LỤC, Word TOC field)
  - numbered body sections + appendices
  - bordered tables, shaded code blocks, callout blockquotes
  - Times New Roman, A4, 1.5 line spacing, page numbers in footer

Run from repo root:
    python docs/build_report_docx.py
"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
MD = ROOT / "docs" / "REPORT_KMS_TU_VI.md"
OUT = ROOT / "docs" / "REPORT_KMS_TU_VI.docx"

BODY_FONT = "Times New Roman"
BODY_SIZE = 12
MONO_FONT = "Consolas"
GRAY = RGBColor(0x55, 0x55, 0x55)

INLINE_RE = re.compile(
    r"(\*\*.*?\*\*|\*[^*]+?\*|`[^`]+?`|\[[^\]]+\]\([^)]+\))"
)


# ------------------------------------------------------------------ helpers
def _add_plain(p, text, font=BODY_FONT, size=BODY_SIZE, bold=False,
               italic=False, code=False, color=None):
    r = p.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    if color is not None:
        r.font.color.rgb = color
    if code:
        r.font.name = MONO_FONT
        r.font.size = Pt(min(size, 10.5))
        _shade_run(r, "F2F2F2")
    return r


def add_inline(p, text, size=BODY_SIZE, base_bold=False, base_italic=False,
               color=None):
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            _add_plain(p, text[pos:m.start()], size=size, bold=base_bold,
                       italic=base_italic, color=color)
        tok = m.group(0)
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            _add_plain(p, tok[2:-2], size=size, bold=True, italic=base_italic,
                       color=color)
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            _add_plain(p, tok[1:-1], size=size, code=True, color=color)
        elif tok.startswith("["):
            inner = tok[1:tok.index("]")]
            _add_plain(p, inner, size=size, bold=base_bold,
                       italic=base_italic, color=color)
        else:  # *italic*
            _add_plain(p, tok[1:-1], size=size, bold=base_bold, italic=True,
                       color=color)
        pos = m.end()
    if pos < len(text):
        _add_plain(p, text[pos:], size=size, bold=base_bold,
                   italic=base_italic, color=color)


def _shade_run(run, fill):
    rPr = run._r.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    rPr.append(shd)


def _shade_paragraph(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def _shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _left_border(p, color="8EAADB", sz=18):
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(sz))
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    pBdr.append(left)
    pPr.append(pBdr)


# ------------------------------------------------------------------ blocks
def add_heading(doc, text, level):
    h = doc.add_heading("", level=level)
    add_inline(h, text, size=BODY_SIZE)
    return h


def add_code_block(doc, code_lines):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(4)
    pf.space_after = Pt(8)
    pf.left_indent = Cm(0.4)
    pf.right_indent = Cm(0.2)
    pf.line_spacing = 1.0
    _shade_paragraph(p, "F5F5F5")
    for i, line in enumerate(code_lines):
        if i > 0:
            p.add_run().add_break()
        _add_plain(p, line, font=MONO_FONT, size=8)


def add_blockquote(doc, quote_lines):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.6)
    pf.right_indent = Cm(0.3)
    pf.space_before = Pt(8)
    pf.space_after = Pt(8)
    _left_border(p, "8EAADB", 18)
    _shade_paragraph(p, "EEF3FA")
    add_inline(p, " ".join(quote_lines), size=BODY_SIZE)


def add_list_item(doc, content, marker, indent):
    level = len(indent) // 2
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.8 + level * 0.75)
    pf.first_line_indent = Cm(-0.6)
    pf.space_after = Pt(2)
    if marker in ("-", "*"):
        mark = ["•", "–", "·"][min(level, 2)]
        _add_plain(p, mark + "   ")
    else:
        _add_plain(p, marker + "  ", bold=True)
    add_inline(p, content)


_SEP_RE = re.compile(r"^:?-{2,}:?$")


def add_md_table(doc, rows):
    parsed = []
    for row in rows:
        cells = [c.strip() for c in row.strip().strip("|").split("|")]
        if all(_SEP_RE.fullmatch(c or "") for c in cells):
            continue
        parsed.append(cells)
    if not parsed:
        return
    ncol = len(parsed[0])
    tbl = doc.add_table(rows=1 + len(parsed) - 1, cols=ncol)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = True
    header, data = parsed[0], parsed[1:]
    for j in range(ncol):
        cell = tbl.cell(0, j)
        cell.text = ""
        _shade_cell(cell, "D9E1F2")
        add_inline(cell.paragraphs[0], header[j] if j < len(header) else "",
                   size=10.5, base_bold=True)
    for i, dr in enumerate(data):
        for j in range(ncol):
            cell = tbl.cell(i + 1, j)
            cell.text = ""
            add_inline(cell.paragraphs[0], dr[j] if j < len(dr) else "",
                       size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_toc_field(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t")
    t.text = "(Bấm F9 / Update Fields trong Word để sinh mục lục)"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (fld_begin, instr, sep, t, end):
        run._r.append(el)


def add_page_number(doc):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.font.size = Pt(10)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for el in (fld_begin, instr, fld_end):
        run._r.append(el)


def _center(doc, text, size=12, bold=False, space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    _add_plain(p, text, size=size, bold=bold)
    return p


def add_cover(doc):
    _center(doc, "BỘ GIÁO DỤC VÀ ĐÀO TẠO", 14, bold=True, space_after=2)
    _center(doc, "HỌC VIỆN [điền tên trường / học viện]", 13, bold=True,
            space_after=2)
    _center(doc, "KHOA / NGÀNH [điền]", 12, space_after=0)
    for _ in range(4):
        doc.add_paragraph()
    _center(doc, "BÁO CÁO DỰ ÁN MÔN HỌC", 16, bold=True, space_after=4)
    _center(doc, "QUẢN TRỊ TRI THỨC (KMS)", 15, bold=True, space_after=16)
    _center(doc, "Hệ thống Quản trị Tri thức (KMS) về Tử Vi,", 14, bold=True,
            space_after=2)
    _center(doc, "có tầng ứng dụng LLM được kiểm soát trích dẫn", 13, bold=True,
            space_after=24)
    for label, val in [
        ("Sinh viên / Nhóm", "[điền tên các thành viên]"),
        ("Giảng viên phụ trách", "[điền tên GV]"),
        ("Học phần", "Quản trị Tri thức (KMS) — [điền mã học phần / tín chỉ]"),
        ("Ngày báo cáo", "01/11/2026"),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        _add_plain(p, label + ": ", size=12, bold=True)
        _add_plain(p, val, size=12)
    doc.add_page_break()


# ------------------------------------------------------------------ setup
def setup_doc(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin, sec.right_margin = Cm(3), Cm(2)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_SIZE)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    for lvl, sz in [(1, 14), (2, 13), (3, 12)]:
        h = doc.styles[f"Heading {lvl}"]
        h.font.name = BODY_FONT
        h.font.size = Pt(sz)
        h.font.bold = True
        h.font.color.rgb = RGBColor(0, 0, 0)
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.line_spacing = 1.3

    cp = doc.core_properties
    cp.title = "Báo cáo dự án KMS Tử Vi"
    cp.author = "Nhóm sinh viên"
    add_page_number(doc)


# ------------------------------------------------------------------ main
LIST_RE = re.compile(r"^(\s*)([-*]|\d+\.)\s+(.*)$")


def build():
    text = MD.read_text(encoding="utf-8")
    lines = text.split("\n")

    doc = Document()
    setup_doc(doc)
    add_cover(doc)

    i = 0
    state = "skip_head"   # until we reach "## Tóm tắt"
    in_toc = False
    while i < len(lines):
        line = lines[i]
        s = line.strip()

        if state == "skip_head":
            if s.startswith("## Tóm tắt"):
                state = "body"
                add_heading(doc, s[3:], 1)
                i += 1
                continue
            i += 1
            continue

        if in_toc:
            if s.startswith("## 1."):
                in_toc = False
            else:
                i += 1
                continue

        if s.startswith("## MỤC LỤC"):
            doc.add_page_break()
            add_heading(doc, "MỤC LỤC", 1)
            add_toc_field(doc)
            doc.add_page_break()
            in_toc = True
            i += 1
            continue

        if s.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1  # closing fence
            add_code_block(doc, block)
            continue

        if s.startswith("|"):
            tbl = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                tbl.append(lines[i])
                i += 1
            add_md_table(doc, tbl)
            continue

        if s.startswith("#### "):
            add_heading(doc, s[5:], 3); i += 1; continue
        if s.startswith("### "):
            add_heading(doc, s[4:], 2); i += 1; continue
        if s.startswith("## "):
            add_heading(doc, s[3:], 1); i += 1; continue
        if s.startswith("# "):
            add_heading(doc, s[2:], 1); i += 1; continue

        if s.startswith(">"):
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip().lstrip(">").strip())
                i += 1
            add_blockquote(doc, quote)
            continue

        m = LIST_RE.match(line)
        if m:
            indent, marker, content = m.group(1), m.group(2), m.group(3)
            add_list_item(doc, content, marker, indent)
            i += 1
            continue

        if re.fullmatch(r"-{3,}", s):
            i += 1
            continue
        if s == "":
            i += 1
            continue

        p = doc.add_paragraph()
        add_inline(p, line)
        i += 1

    doc.save(OUT)
    print("WROTE", OUT)


if __name__ == "__main__":
    build()
